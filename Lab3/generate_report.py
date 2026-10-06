import json
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_screenshot_placeholder(doc, title="[CHÈN HÌNH ẢNH MINH CHỨNG TẠI ĐÂY]", lines=5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.font.color.rgb = RGBColor(255, 0, 0)
    run.italic = True
    for _ in range(lines):
        doc.add_paragraph()

def create_report():
    # Read evaluation results
    json_path = "evaluation_results.json"
    if os.path.exists(json_path):
        with open(json_path, "r", encoding="utf-8") as f:
            eval_data = json.load(f)
    else:
        eval_data = {"summary": [], "details": []}

    doc = Document()
    
    # Title
    title = doc.add_heading("BÁO CÁO CHI TIẾT BÀI THỰC HÀNH #3:\nDỰNG AGENT ĐẶT VÉ MÁY BAY", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("Môn học: SE373 - Tác tử trí tuệ nhân tạo (Agentic AI)").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("Họ và tên: Đinh Nguyễn Đức Tâm - 23521384").alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # Mở đầu
    doc.add_heading("I. Giới thiệu và Cấu trúc dự án", level=1)
    doc.add_paragraph("Trong bài thực hành này, mục tiêu là xây dựng các Agent có khả năng hỗ trợ người dùng đặt vé máy bay một cách tự động, đồng thời đảm bảo an toàn thông qua các cơ chế rào cản (harness). Báo cáo này trình bày chi tiết về các thành phần đã xây dựng, các mẫu thiết kế Agent được áp dụng và kết quả đánh giá.")
    
    doc.add_heading("Cấu trúc thư mục mã nguồn:", level=2)
    doc.add_paragraph("Dưới đây là cấu trúc của dự án. Sinh viên thực hiện lệnh `tree` hoặc `ls -R` để hiển thị danh sách các file và thư mục.")
    add_screenshot_placeholder(doc, "[CHÈN HÌNH ẢNH MÀN HÌNH CẤU TRÚC THƯ MỤC / LỆNH TREE TẠI ĐÂY]")

    # Section 1
    doc.add_heading("II. Cài đặt các lớp Harness bảo vệ", level=1)
    p = doc.add_paragraph()
    p.add_run("Các lớp harness được thiết kế trong file ")
    p.add_run("lib/harness.py").italic = True
    p.add_run(" nhằm kiểm soát luồng thực thi, đảm bảo an toàn và đúng đắn. Một Agent tự trị cần được kiểm soát chặt chẽ để tránh các hành vi ngoài ý muốn, đặc biệt là trong các hệ thống có tác động đến tài chính như đặt vé máy bay.")
    
    doc.add_heading("1. Ràng buộc dữ liệu (BookingConstraints)", level=2)
    doc.add_paragraph("Lớp BookingConstraints đảm bảo thông tin vé (điểm đi, điểm đến, ngày khởi hành, khung giờ, trần giá) luôn được tuân thủ nghiêm ngặt. Nếu Agent tìm hoặc đặt vé vượt mức trần giá cho phép (ví dụ 1.500.000 VNĐ), hoặc sai định dạng ngày tháng, harness sẽ phát hiện và từ chối hành động ngay lập tức.")
    add_screenshot_placeholder(doc, "[CHÈN HÌNH ẢNH TEST/CHẠY LỆNH KIỂM TRA BOOKING CONSTRAINTS TẠI ĐÂY]")

    doc.add_heading("2. Tiêu chí hoàn thành (CompletionVerifier)", level=2)
    doc.add_paragraph("Lớp CompletionVerifier đóng vai trò như một màng lọc cuối cùng trước khi Agent trả về kết quả cho người dùng. Nó kiểm tra tính đầy đủ của quá trình: Agent chỉ được coi là hoàn thành nhiệm vụ khi đã có mã vé (Ticket ID), trạng thái vé là BOOKED, và tên hành khách/CCCD phải khớp hoàn toàn với yêu cầu ban đầu của người dùng.")
    
    doc.add_heading("3. Kiểm soát quyền & Hạn mức (PermissionGuard)", level=2)
    doc.add_paragraph("Lớp PermissionGuard chịu trách nhiệm giám sát các hành động nhạy cảm như thanh toán trừ tiền trực tiếp hoặc đặt các loại vé không hoàn/hủy. Khi số tiền vượt quá một ngưỡng an toàn, hành động sẽ bị tạm dừng và hệ thống yêu cầu sự phê duyệt từ con người (Human-in-the-loop) với trạng thái AWAITING_APPROVAL.")
    add_screenshot_placeholder(doc, "[CHÈN HÌNH ẢNH MINH CHỨNG PERMISSION GUARD CHẶN HÀNH ĐỘNG/YÊU CẦU XÁC NHẬN TẠI ĐÂY]")

    doc.add_heading("4. Quản lý bàn giao (HandoffManager)", level=2)
    doc.add_paragraph("Lớp HandoffManager sẽ tự động ngắt và chuyển quyền xử lý cho con người (hoặc hệ thống fallback) khi phát hiện các biểu hiện bất thường như: Agent bị kẹt trong vòng lặp lỗi (Loop inducement), liên tục vi phạm trần giá, hoặc có dấu hiệu tạo ra dữ kiện ảo (hallucination).")

    # Section 2
    doc.add_heading("III. Cài đặt Agent với 3 mẫu thiết kế (Design Patterns)", level=1)
    
    doc.add_heading("1. Mẫu ReAct (Reasoning & Acting)", level=2)
    doc.add_paragraph("Được cài đặt tại lib/agent_react.py. Mô hình hoạt động theo chu trình Suy nghĩ (Thought) -> Hành động (Action) -> Quan sát (Observation). Mô hình này rất phù hợp với các công việc cần linh hoạt điều chỉnh bước đi dựa trên phản hồi của công cụ. Tuy nhiên, nhược điểm là dễ bị kẹt trong vòng lặp vô tận nếu công cụ liên tục trả về lỗi.")
    doc.add_paragraph("Minh chứng chạy thử Agent ReAct:")
    add_screenshot_placeholder(doc, "[CHÈN HÌNH ẢNH TERMINAL KHI CHẠY AGENT REACT TẠI ĐÂY]")

    doc.add_heading("2. Mẫu Plan-then-Execute", level=2)
    doc.add_paragraph("Được cài đặt tại lib/agent_plan_execute.py. Mô hình này tách biệt rõ ràng 2 giai đoạn: Lập kế hoạch (Planner) và Thực thi (Executor). Giúp luồng đi mạch lạc và dự đoán được toàn bộ các bước trước khi thực sự gọi bất kỳ công cụ nào. Nhược điểm của phương pháp này là độ trễ thường cao do LLM phải suy nghĩ sinh ra kế hoạch chi tiết từ đầu, và thiếu tính linh hoạt khi xảy ra lỗi đột xuất ở giữa chừng.")
    doc.add_paragraph("Minh chứng chạy thử Agent Plan-then-Execute:")
    add_screenshot_placeholder(doc, "[CHÈN HÌNH ẢNH TERMINAL KHI CHẠY AGENT PLAN-THEN-EXECUTE TẠI ĐÂY]")

    doc.add_heading("3. Mẫu Lai (Hybrid)", level=2)
    doc.add_paragraph("Được cài đặt tại lib/agent_hybrid.py. Đây là sự kết hợp hoàn hảo ưu điểm của cả hai mô hình trên: Vừa có một bản kế hoạch sơ bộ định hướng luồng công việc từ ban đầu, vừa giữ được sự linh hoạt của vòng lặp ReAct trong quá trình thực thi từng tác vụ nhỏ. Thiết kế này mang lại hiệu năng ổn định và tỷ lệ thành công cao nhất, đặc biệt trong các tình huống thực tế phức tạp.")
    doc.add_paragraph("Minh chứng chạy thử Agent Hybrid:")
    add_screenshot_placeholder(doc, "[CHÈN HÌNH ẢNH TERMINAL KHI CHẠY AGENT HYBRID TẠI ĐÂY]")

    # Section 3
    doc.add_heading("IV. Đánh giá hiệu năng và Kết quả (Evaluation)", level=1)
    doc.add_paragraph("Quá trình đánh giá được thực hiện thông qua script tự động (run_evaluation.py). Script này chạy 3 Agent qua 4 kịch bản chuẩn hóa:")
    doc.add_paragraph("- Kịch bản 1: Happy path (Mọi thứ thuận lợi).")
    doc.add_paragraph("- Kịch bản 2: Over budget (Giá vé vượt ngân sách).")
    doc.add_paragraph("- Kịch bản 3: Loop inducement (Mô phỏng API lỗi để bẫy vòng lặp).")
    doc.add_paragraph("- Kịch bản 4: Permission check (Yêu cầu thanh toán cần xét duyệt).")
    doc.add_paragraph("Dưới đây là bảng tổng hợp kết quả:")

    summary_data = eval_data.get("summary", [])
    if summary_data:
        table = doc.add_table(rows=1, cols=7)
        table.style = 'Table Grid'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Mẫu thiết kế'
        hdr_cells[1].text = 'Tỷ lệ thành công (%)'
        hdr_cells[2].text = 'Số bước TB'
        hdr_cells[3].text = 'Token TB'
        hdr_cells[4].text = 'Độ trễ TB (ms)'
        hdr_cells[5].text = 'Bàn giao an toàn'
        hdr_cells[6].text = 'Grounding (%)'

        for row in summary_data:
            row_cells = table.add_row().cells
            row_cells[0].text = str(row.get("pattern", ""))
            row_cells[1].text = str(row.get("success_rate_pct", ""))
            row_cells[2].text = str(row.get("avg_steps", ""))
            row_cells[3].text = str(row.get("avg_tokens", ""))
            row_cells[4].text = str(row.get("avg_latency_ms", ""))
            row_cells[5].text = f"{row.get('safe_handoff_count', 0)}/4"
            row_cells[6].text = str(row.get("grounding_accuracy_pct", ""))
    
    doc.add_paragraph("\nMinh chứng chạy đánh giá (run_evaluation.py):")
    add_screenshot_placeholder(doc, "[CHÈN HÌNH ẢNH TERMINAL KẾT QUẢ CHẠY RUN_EVALUATION.PY TẠI ĐÂY]")

    doc.add_heading("Phân tích kết quả đánh giá:", level=2)
    doc.add_paragraph("- ReAct: Thể hiện tính linh hoạt cao, tuy nhiên số bước trung bình đôi khi lớn do đặc tính phải thăm dò liên tục. Khi gặp công cụ lỗi (kịch bản 3), ReAct dễ bị cuốn vào vòng lặp retry nếu không có HandoffManager can thiệp kịp thời.")
    doc.add_paragraph("- Plan-then-Execute: Giữ được luồng thực thi rất ổn định, hầu như không bị loop. Đổi lại, số lượng Token tiêu thụ và độ trễ phản hồi (latency) cao hơn rõ rệt vì LLM phải tổng hợp thông tin và sinh ra kế hoạch dài ngay từ đầu. Khi thực thi thất bại, việc điều chỉnh lại toàn bộ kế hoạch khá khó khăn.")
    doc.add_paragraph("- Lai (Hybrid): Đạt được sự cân bằng tối ưu nhất giữa tốc độ và độ tin cậy. Số lượng token vừa phải, thời gian phản hồi tốt, và đặc biệt tỷ lệ kích hoạt bàn giao an toàn (safe handoff) hoạt động chính xác nhất.")

    # Section 4
    doc.add_heading("V. Kết luận", level=1)
    doc.add_paragraph("Bài thực hành đã hoàn thành việc xây dựng và so sánh 3 mẫu thiết kế Agent phổ biến, kết hợp với các cơ chế bảo vệ chặt chẽ (Harness). Việc áp dụng Harness không chỉ giúp Agent tuân thủ các quy tắc nghiệp vụ mà còn ngăn chặn rủi ro phát sinh trong thực tế. Mẫu Hybrid chứng minh là lựa chọn phù hợp nhất để triển khai cho các Agent có độ phức tạp trung bình đến cao.")
    
    # Save
    doc.save("23521384_DinhNguyenDucTam_BTTH3.docx")
    print("Report generated successfully as 23521384_DinhNguyenDucTam_BTTH3.docx")

if __name__ == '__main__':
    create_report()

